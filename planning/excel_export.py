"""Export one SKU's full planning calculation as a live-formula Excel workbook.

The workbook mirrors the structure of the Excel validation model, but contains only the selected
SKU. Every calculated cell is a real Excel formula (walk-forward forecasts, WAPE, bias, error SD,
method selection, z-value, safety stock, reorder point, inventory position, EOQ, order decision,
12-week projection and status), so a planner can change an input and watch the chain recalculate.
A "Check" sheet compares each Excel result with the Python engine's value.
"""

from __future__ import annotations

import io
from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

from planning.forecasting import METHOD_LABELS, METHODS
from planning.pipeline import PlanResult

# ------------------------------------------------------------------ styles (same conventions as the validation model)
INK = "111633"
INDIGO = "4F46E5"
HEADER_FILL = PatternFill("solid", fgColor=INK)
SECTION_FILL = PatternFill("solid", fgColor="EEF0FF")
INPUT_FILL = PatternFill("solid", fgColor="FFF7D6")
RESULT_FILL = PatternFill("solid", fgColor="E8F7EF")
WHITE_BOLD = Font(bold=True, color="FFFFFF")
TITLE = Font(bold=True, size=16, color=INK)
SUBTITLE = Font(italic=True, size=10, color="5A6178")
BOLD = Font(bold=True, color="151A30")
INPUT_FONT = Font(color="1D4ED8")
LINK_FONT = Font(color="047857")
THIN = Side(style="thin", color="E4E7F0")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")

HISTORY_FIRST_ROW = 5  # Forecast sheet: week 1 sits on this row


def _header(ws, row: int, labels: list[str], widths: list[float] | None = None) -> None:
    for i, label in enumerate(labels, start=1):
        c = ws.cell(row, i, label)
        c.fill = HEADER_FILL
        c.font = WHITE_BOLD
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = BOX
    if widths:
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w


def _title(ws, title: str, subtitle: str) -> None:
    ws["A1"] = title
    ws["A1"].font = TITLE
    ws["A2"] = subtitle
    ws["A2"].font = SUBTITLE


def _name(wb: Workbook, name: str, ref: str) -> None:
    wb.defined_names[name] = DefinedName(name, attr_text=ref)


def build_sku_workbook(result: PlanResult, sku: str) -> Workbook:
    """Build the workbook in memory (see :func:`sku_workbook_bytes` for the downloadable file)."""
    data, params = result.data, result.params
    row = result.sku_plan(sku)
    fc = result.forecasts[sku]
    proj = result.projections[sku]
    pos = data.open_orders(sku)
    history = fc.walk_forward
    n_weeks = len(history)
    eval_start, eval_end = fc.evaluation_window
    scenario = data.scenario

    wb = Workbook()
    wb.calculation.fullCalcOnLoad = True

    # ================================================================== Summary (filled at the end)
    summary = wb.active
    summary.title = "Summary"

    # ================================================================== Inputs
    inp = wb.create_sheet("Inputs")
    _title(
        inp,
        f"Inputs · {sku} {row['product_name']}",
        "Blue text on yellow = input you may change. Everything else is a formula.",
    )
    _header(inp, 4, ["Parameter", "Value", "Unit", "Note"], [34, 16, 16, 70])
    inputs = [
        (
            "TSL",
            "Target service level",
            params.target_service_level,
            "%",
            "Cycle service level used for the z-value.",
            "0.0%",
        ),
        (
            "Holding_Rate",
            "Annual holding rate",
            params.annual_holding_rate,
            "% of cost / year",
            "Capital, storage, obsolescence.",
            "0%",
        ),
        (
            "Weeks_Per_Year",
            "Weeks per year",
            params.weeks_per_year,
            "weeks",
            "Annualises the weekly forecast for EOQ.",
            "0",
        ),
        (
            "Excess_WoC",
            "Excess-stock threshold",
            params.excess_weeks_of_cover,
            "weeks of cover",
            "Cover above this is EXCESS STOCK.",
            "0",
        ),
        ("SES_Alpha", "SES smoothing alpha", params.ses_alpha, "0-1", "Weight on the latest actual.", "0.00"),
        (
            "WAPE_Tol",
            "WAPE tie tolerance",
            params.wape_tolerance,
            "WAPE points",
            "A simpler method wins if within this of the best.",
            "0.0%",
        ),
        (
            "Lead_Time",
            "Supplier lead time",
            int(row["lead_time_weeks"]),
            "weeks",
            f"Supplier {row['supplier_id']} · {row['supplier_name']} (fixed).",
            "0",
        ),
        (
            "Ordering_Cost",
            "Ordering cost per order",
            float(row["ordering_cost_eur"]),
            "EUR / order",
            "Fixed cost of placing and receiving an order.",
            "€#,##0.00",
        ),
        ("Unit_Cost", "Unit cost", float(row["unit_cost_eur"]), "EUR / unit", "", "€#,##0.00"),
        (
            "Case_Pack",
            "Case pack",
            int(row["case_pack"]),
            "units",
            "Orders are rounded up to full cases.",
            "0",
        ),
        (
            "On_Hand",
            "On hand",
            int(row["on_hand"]),
            "units",
            f"At the snapshot date {scenario.snapshot_date:%d %b %Y}.",
            "#,##0",
        ),
        (
            "Backorders",
            "Backorders",
            int(row["backorders"]),
            "units",
            "Demand already owed to customers.",
            "#,##0",
        ),
    ]
    r = 5
    for name, label, value, unit, note, fmt in inputs:
        inp.cell(r, 1, label).border = BOX
        c = inp.cell(r, 2, value)
        c.font, c.fill, c.border, c.number_format = INPUT_FONT, INPUT_FILL, BOX, fmt
        inp.cell(r, 3, unit).border = BOX
        inp.cell(r, 4, note).border = BOX
        _name(wb, name, f"Inputs!$B${r}")
        r += 1
    z_row = r
    inp.cell(r, 1, "z-value").border = BOX
    inp.cell(r, 2, "=_xlfn.NORM.S.INV(TSL)").number_format = "0.0000"
    inp.cell(r, 2).border = BOX
    inp.cell(r, 3, "number").border = BOX
    inp.cell(r, 4, "Formula: NORM.S.INV(target service level). 95% ≈ 1.645.").border = BOX
    _name(wb, "Z_Value", f"Inputs!$B${z_row}")

    r += 2
    inp.cell(r, 1, "OPEN PURCHASE ORDERS").font = BOLD
    r += 1
    _header(inp, r, ["PO", "Quantity (units)", "Due week (1 = snapshot week)", "Note"])
    po_first = r + 1
    po_rows = max(len(pos), 1) + 2  # spare rows so a planner can add a PO
    for i in range(po_rows):
        rr = po_first + i
        po = pos.iloc[i] if i < len(pos) else None
        inp.cell(rr, 1, po["po_id"] if po is not None else "").border = BOX
        for col, val in (
            (2, int(po["quantity"]) if po is not None else None),
            (3, int(po["due_week"]) if po is not None else None),
        ):
            c = inp.cell(rr, col, val)
            c.font, c.fill, c.border = INPUT_FONT, INPUT_FILL, BOX
        inp.cell(rr, 4, "" if po is not None else "spare row: add a PO here").border = BOX
    po_last = po_first + po_rows - 1
    _name(wb, "PO_Qty", f"Inputs!$B${po_first}:$B${po_last}")
    _name(wb, "PO_Week", f"Inputs!$C${po_first}:$C${po_last}")
    r = po_last + 1
    inp.cell(r, 1, "On order (sum of open POs)").border = BOX
    inp.cell(r, 2, "=SUM(PO_Qty)").border = BOX
    inp.cell(r, 2).number_format = "#,##0"
    inp.cell(r, 4, "Formula: SUM of the PO quantities above.").border = BOX
    _name(wb, "On_Order", f"Inputs!$B${r}")
    inp.freeze_panes = "A5"

    # ================================================================== Forecast (walk-forward)
    fs = wb.create_sheet("Forecast")
    _title(
        fs,
        "Walk-forward forecasts and errors",
        f"Each forecast uses only earlier weeks. Error = forecast − actual. Evaluation window: weeks {eval_start}–{eval_end} (shaded).",
    )
    _header(
        fs,
        4,
        [
            "Week",
            "Week start",
            "Actual",
            "Naive",
            "MA4",
            "MA8",
            "SES",
            "Err Naive",
            "Err MA4",
            "Err MA8",
            "Err SES",
            "|Err| Naive",
            "|Err| MA4",
            "|Err| MA8",
            "|Err| SES",
            "Selected forecast",
        ],
        [8, 12, 10, 10, 10, 10, 10, 11, 11, 11, 11, 11, 11, 11, 11, 13],
    )
    first = HISTORY_FIRST_ROW
    last = first + n_weeks - 1
    for i, (week, actual) in enumerate(history["actual"].items()):
        rr = first + i
        week = int(week)
        fs.cell(rr, 1, week)
        fs.cell(
            rr, 2, scenario.snapshot_date - timedelta(weeks=n_weeks - week + 1)
        ).number_format = "dd mmm yy"
        c = fs.cell(rr, 3, float(actual))
        c.font, c.fill = INPUT_FONT, INPUT_FILL
        if week >= 2:
            fs.cell(rr, 4, f"=C{rr - 1}")
            fs.cell(rr, 7, f"=C{rr - 1}" if week == 2 else f"=SES_Alpha*C{rr - 1}+(1-SES_Alpha)*G{rr - 1}")
        if week >= 5:
            fs.cell(rr, 5, f"=AVERAGE(C{rr - 4}:C{rr - 1})")
        if week >= 9:
            fs.cell(rr, 6, f"=AVERAGE(C{rr - 8}:C{rr - 1})")
        for j, col in enumerate("DEFG"):
            min_week = {"D": 2, "E": 5, "F": 9, "G": 2}[col]
            if week >= min_week:
                fs.cell(rr, 8 + j, f"={col}{rr}-C{rr}")
                fs.cell(rr, 12 + j, f"=ABS({get_column_letter(8 + j)}{rr})")
        if week >= 9:
            fs.cell(rr, 16, f"=INDEX(D{rr}:G{rr},Selected_No)")
        for col in range(4, 17):
            fs.cell(rr, col).number_format = "#,##0.0"
        if eval_start <= week <= eval_end:
            for col in range(1, 17):
                if col != 3:
                    fs.cell(rr, col).fill = SECTION_FILL
    ev_first, ev_last = first + eval_start - 1, first + eval_end - 1
    fs.freeze_panes = "C5"

    # ================================================================== Method selection
    ms = wb.create_sheet("Method_Selection")
    _title(
        ms,
        "Forecast method comparison and selection",
        f"Scored on weeks {eval_start}–{eval_end} (Forecast rows {ev_first}–{ev_last}). Lowest WAPE wins; a simpler method wins if within the tolerance.",
    )
    _header(
        ms,
        4,
        ["Method", "WAPE", "Bias", "Error SD", "Next-week forecast", "Within tolerance?", "Formula notes"],
        [30, 11, 11, 11, 18, 17, 70],
    )
    err_cols = {"Naive": ("H", "L"), "MA4": ("I", "M"), "MA8": ("J", "N"), "SES": ("K", "O")}
    next_formula = {
        "Naive": f"=Forecast!C{last}",
        "MA4": f"=AVERAGE(Forecast!C{last - 3}:C{last})",
        "MA8": f"=AVERAGE(Forecast!C{last - 7}:C{last})",
        "SES": f"=SES_Alpha*Forecast!C{last}+(1-SES_Alpha)*Forecast!G{last}",
    }
    actual_sum = f"SUM(Forecast!C{ev_first}:C{ev_last})"
    for i, m in enumerate(METHODS):
        rr = 5 + i
        err, abs_err = err_cols[m]
        ms.cell(rr, 1, f"{m} · {METHOD_LABELS[m]}")
        ms.cell(
            rr, 2, f"=SUM(Forecast!{abs_err}{ev_first}:{abs_err}{ev_last})/{actual_sum}"
        ).number_format = "0.0%"
        ms.cell(
            rr, 3, f"=SUM(Forecast!{err}{ev_first}:{err}{ev_last})/{actual_sum}"
        ).number_format = "+0.0%;-0.0%"
        ms.cell(rr, 4, f"=_xlfn.STDEV.S(Forecast!{err}{ev_first}:{err}{ev_last})").number_format = "0.00"
        ms.cell(rr, 5, next_formula[m]).number_format = "#,##0.00"
        ms.cell(rr, 6, f"=B{rr}<=Best_WAPE+WAPE_Tol")
        ms.cell(rr, 7, "WAPE = Σ|F−A| ÷ ΣA · Bias = Σ(F−A) ÷ ΣA · SD = STDEV.S(errors)" if i == 0 else "")
        for col in range(1, 8):
            ms.cell(rr, col).border = BOX
    ms.cell(10, 1, "Best WAPE").font = BOLD
    ms.cell(10, 2, "=MIN(B5:B8)").number_format = "0.0%"
    _name(wb, "Best_WAPE", "Method_Selection!$B$10")
    ms.cell(11, 1, "Selected method # (first within tolerance)").font = BOLD
    ms.cell(11, 2, "=MATCH(TRUE,F5:F8,0)")
    _name(wb, "Selected_No", "Method_Selection!$B$11")
    selected_rows = [
        ("Selected method", '=CHOOSE(Selected_No,"Naive","MA4","MA8","SES")', "@", "Sel_Method"),
        ("Weekly forecast (units)", "=INDEX(E5:E8,Selected_No)", "#,##0.00", "Forecast"),
        ("WAPE", "=INDEX(B5:B8,Selected_No)", "0.0%", "Sel_WAPE"),
        ("Bias", "=INDEX(C5:C8,Selected_No)", "+0.0%;-0.0%", "Sel_Bias"),
        ("Forecast error SD (units)", "=INDEX(D5:D8,Selected_No)", "0.00", "Error_SD"),
    ]
    for i, (label, formula, fmt, name) in enumerate(selected_rows):
        rr = 12 + i
        ms.cell(rr, 1, label).font = BOLD
        c = ms.cell(rr, 2, formula)
        c.number_format, c.fill, c.border = fmt, RESULT_FILL, BOX
        _name(wb, name, f"Method_Selection!$B${rr}")

    # ================================================================== Replenishment chain
    rp = wb.create_sheet("Replenishment")
    _title(
        rp,
        "Safety stock → reorder point → inventory position → EOQ → order",
        "ROP answers WHEN to order; EOQ answers approximately HOW MUCH.",
    )
    _header(
        rp,
        4,
        ["#", "Calculation", "Formula in words", "Result", "Unit", "What it means"],
        [5, 30, 44, 14, 14, 70],
    )
    chain = [
        (
            "Safety_Stock",
            "Safety stock",
            "z × error SD × √lead time",
            "=Z_Value*Error_SD*SQRT(Lead_Time)",
            "units",
            "Buffer against forecast error over the lead time.",
            "#,##0.00",
        ),
        (
            "LT_Demand",
            "Lead-time demand",
            "weekly forecast × lead time",
            "=Forecast*Lead_Time",
            "units",
            "Expected demand while a new order is on its way.",
            "#,##0.00",
        ),
        (
            "ROP",
            "Reorder point",
            "lead-time demand + safety stock",
            "=LT_Demand+Safety_Stock",
            "units",
            "Order when the inventory position is at or below this.",
            "#,##0.00",
        ),
        (
            "Inv_Position",
            "Inventory position",
            "on hand + on order − backorders",
            "=On_Hand+On_Order-Backorders",
            "units",
            "Counting stock on order prevents a duplicate order.",
            "#,##0",
        ),
        (
            "WoC",
            "Weeks of cover",
            "max(position, 0) ÷ weekly forecast",
            "=IF(Forecast>0,MAX(Inv_Position,0)/Forecast,0)",
            "weeks",
            "Negative positions show 0 cover; the shortage is explicit.",
            "0.0",
        ),
        (
            "Annual_Demand",
            "Annual forecast demand",
            "weekly forecast × 52",
            "=Forecast*Weeks_Per_Year",
            "units / year",
            "EOQ works on annual quantities.",
            "#,##0.00",
        ),
        (
            "Holding_Cost",
            "Annual holding cost per unit",
            "unit cost × holding rate",
            "=Unit_Cost*Holding_Rate",
            "EUR / unit / yr",
            "Cost of keeping one unit in stock for a year.",
            "€#,##0.00",
        ),
        (
            "EOQ",
            "Economic order quantity",
            "√(2 × annual demand × ordering cost ÷ holding cost)",
            "=IF(Annual_Demand>0,SQRT(2*Annual_Demand*Ordering_Cost/Holding_Cost),0)",
            "units",
            "Balances ordering cost against holding cost.",
            "#,##0.00",
        ),
        (
            "EOQ_Rounded",
            "EOQ rounded to case pack",
            "ROUNDUP(EOQ ÷ case pack) × case pack",
            "=ROUNDUP(EOQ/Case_Pack,0)*Case_Pack",
            "units",
            "Suppliers ship full cases.",
            "#,##0",
        ),
        (
            "Trigger",
            "Reorder trigger",
            "is inventory position ≤ reorder point?",
            '=IF(Inv_Position<=ROP,"YES","NO")',
            "",
            "YES means an order is placed at this review.",
            "@",
        ),
        (
            "Rec_Order",
            "Recommended order",
            "if triggered: rounded EOQ, else 0",
            '=IF(Trigger="YES",EOQ_Rounded,0)',
            "units",
            "The order proposed this week.",
            "#,##0",
        ),
        (
            "Order_Value",
            "Order value",
            "recommended order × unit cost",
            "=Rec_Order*Unit_Cost",
            "EUR",
            "Cash committed by this order.",
            "€#,##0",
        ),
        (
            "Status",
            "Planning status",
            "stockout risk → reorder → excess → healthy",
            '=IF(Zero_In_LT="YES","STOCKOUT RISK",IF(Trigger="YES","REORDER REQUIRED",IF(WoC>Excess_WoC,"EXCESS STOCK","HEALTHY")))',
            "",
            "Priority order, one status per SKU.",
            "@",
        ),
    ]
    for i, (name, label, words, formula, unit, meaning, fmt) in enumerate(chain):
        rr = 5 + i
        rp.cell(rr, 1, i + 1)
        rp.cell(rr, 2, label).font = BOLD
        rp.cell(rr, 3, words)
        c = rp.cell(rr, 4, formula)
        c.number_format, c.fill = fmt, RESULT_FILL
        rp.cell(rr, 5, unit)
        rp.cell(rr, 6, meaning).alignment = WRAP
        for col in range(1, 7):
            rp.cell(rr, col).border = BOX
        _name(wb, name, f"Replenishment!$D${rr}")

    # ================================================================== Projection
    pj = wb.create_sheet("Projection")
    _title(
        pj,
        "12-week projected inventory",
        "Closing = opening + existing PO receipts + new order receipt − forecast demand. Only today's order is included.",
    )
    _header(
        pj,
        4,
        [
            "Week",
            "Week start",
            "Opening",
            "Existing PO receipt",
            "New order receipt",
            "Forecast demand",
            "Closing",
            "Zero or below?",
        ],
        [8, 12, 12, 16, 16, 15, 12, 14],
    )
    horizon = params.projection_weeks
    for w in range(1, horizon + 1):
        rr = 4 + w
        pj.cell(rr, 1, w)
        pj.cell(rr, 2, scenario.week_start(w)).number_format = "dd mmm yy"
        pj.cell(rr, 3, "=On_Hand-Backorders" if w == 1 else f"=G{rr - 1}")
        pj.cell(rr, 4, f"=SUMIFS(PO_Qty,PO_Week,A{rr})")
        pj.cell(rr, 5, f"=IF(AND(Rec_Order>0,A{rr}=Lead_Time+1),Rec_Order,0)")
        pj.cell(rr, 6, "=Forecast")
        pj.cell(rr, 7, f"=C{rr}+D{rr}+E{rr}-F{rr}")
        pj.cell(rr, 8, f"=IF(G{rr}<=0,1,0)")
        for col in range(3, 8):
            pj.cell(rr, col).number_format = "#,##0.00"
        pj.cell(rr, 7).fill = RESULT_FILL
    p_first, p_last = 5, 4 + horizon
    base = p_last + 2
    pj.cell(base, 1, "First projected zero week").font = BOLD
    pj.cell(base, 4, f'=IFERROR(MATCH(1,H{p_first}:H{p_last},0),"None")').fill = RESULT_FILL
    _name(wb, "First_Zero", f"Projection!$D${base}")
    pj.cell(base + 1, 1, "Zero within lead time?").font = BOLD
    pj.cell(
        base + 1, 4, f'=IF(COUNTIFS(A{p_first}:A{p_last},"<="&Lead_Time,H{p_first}:H{p_last},1)>0,"YES","NO")'
    ).fill = RESULT_FILL
    _name(wb, "Zero_In_LT", f"Projection!$D${base + 1}")
    pj.cell(base + 2, 1, "Lowest projected closing").font = BOLD
    pj.cell(base + 2, 4, f"=MIN(G{p_first}:G{p_last})").fill = RESULT_FILL

    chart = LineChart()
    chart.title, chart.height, chart.width = "Projected closing inventory", 7.5, 16
    chart.y_axis.title, chart.x_axis.title = "Units", "Week"
    chart.add_data(Reference(pj, min_col=7, min_row=4, max_row=p_last), titles_from_data=True)
    chart.set_categories(Reference(pj, min_col=1, min_row=p_first, max_row=p_last))
    pj.add_chart(chart, "J4")

    # ================================================================== Check (Excel formulas vs Python engine)
    ck = wb.create_sheet("Check")
    _title(
        ck,
        "Excel formulas vs the Python engine",
        "Each Excel result is recalculated on opening. PASS means it matches the app to 1e-6 (text must match exactly).",
    )
    _header(
        ck, 4, ["Check item", "Excel result", "Python result", "Difference", "Result"], [34, 16, 16, 14, 10]
    )
    checks = [
        ("Selected method", "=Sel_Method", fc.selected_method),
        ("Weekly forecast", "=Forecast", fc.weekly_forecast),
        ("WAPE", "=Sel_WAPE", fc.wape),
        ("Bias", "=Sel_Bias", fc.bias),
        ("Forecast error SD", "=Error_SD", fc.error_sd),
        ("z-value", "=Z_Value", params.z_value),
        ("Safety stock", "=Safety_Stock", float(row["safety_stock"])),
        ("Lead-time demand", "=LT_Demand", float(row["lead_time_demand"])),
        ("Reorder point", "=ROP", float(row["reorder_point"])),
        ("Inventory position", "=Inv_Position", float(row["inventory_position"])),
        ("Weeks of cover", "=WoC", float(row["weeks_of_cover"])),
        ("Annual demand", "=Annual_Demand", float(row["annual_demand"])),
        ("Holding cost per unit", "=Holding_Cost", float(row["holding_cost_per_unit"])),
        ("EOQ", "=EOQ", float(row["eoq"])),
        ("EOQ rounded", "=EOQ_Rounded", float(row["eoq_rounded"])),
        ("Reorder trigger", "=Trigger", "YES" if row["reorder_triggered"] else "NO"),
        ("Recommended order", "=Rec_Order", float(row["recommended_order"])),
        (
            "First projected zero week",
            "=First_Zero",
            proj.first_zero_week if proj.first_zero_week else "None",
        ),
        (
            "Closing inventory, week 12",
            f"=Projection!G{p_last}",
            float(proj.table["closing_inventory"].iloc[-1]),
        ),
        ("Status", "=Status", row["status"]),
    ]
    for i, (label, formula, py_value) in enumerate(checks):
        rr = 5 + i
        ck.cell(rr, 1, label)
        ck.cell(rr, 2, formula)
        ck.cell(rr, 3, py_value)
        if isinstance(py_value, str):
            ck.cell(rr, 4, "text")
            ck.cell(rr, 5, f'=IF(B{rr}=C{rr},"PASS","FAIL")')
        else:
            ck.cell(rr, 4, f"=ABS(B{rr}-C{rr})").number_format = "0.0E+00"
            ck.cell(rr, 5, f'=IF(D{rr}<=0.000001,"PASS","FAIL")')
            for col in (2, 3):
                ck.cell(rr, col).number_format = "#,##0.0000"
        for col in range(1, 6):
            ck.cell(rr, col).border = BOX
    total = 5 + len(checks)
    ck.cell(total + 1, 1, "Checks passed").font = BOLD
    ck.cell(total + 1, 2, f'=COUNTIF(E5:E{total - 1},"PASS")&" of {len(checks)}"').font = BOLD
    _name(wb, "Checks_Passed", f"Check!$B${total + 1}")

    # ================================================================== Summary
    s = summary
    _title(
        s,
        f"{sku} · {row['product_name']}",
        f"Planning workbook exported {date.today():%d %b %Y} · snapshot {scenario.snapshot_date:%d %b %Y} · {scenario.company} (fictitious company, synthetic data)",
    )
    s.column_dimensions["A"].width = 34
    s.column_dimensions["B"].width = 20
    s.column_dimensions["C"].width = 80
    lines = [
        ("Category", row["category"], None),
        ("Supplier", f"{row['supplier_id']} · {row['supplier_name']}", None),
        ("Status", "=Status", "Stockout risk → reorder required → excess stock → healthy"),
        ("Why (from the app)", row["reason"], None),
        (
            "Selected forecast method",
            "=Sel_Method",
            "Lowest walk-forward WAPE; simpler method wins within tolerance.",
        ),
        ("Weekly forecast (units)", "=Forecast", "Flat for the next 12 weeks."),
        ("Forecast accuracy (1 − WAPE)", "=1-Sel_WAPE", "Out-of-sample, last 52 weeks."),
        ("Safety stock (units)", "=Safety_Stock", "z × error SD × √lead time"),
        ("Reorder point (units)", "=ROP", "WHEN to order"),
        ("Inventory position (units)", "=Inv_Position", "on hand + on order − backorders"),
        ("EOQ (units)", "=EOQ", "HOW MUCH to order"),
        ("Recommended order (units)", "=Rec_Order", "EOQ rounded up to the case pack, if triggered"),
        ("Order value", "=Order_Value", ""),
        ("Validation", "=Checks_Passed", "Excel formulas vs Python engine (see the Check sheet)."),
    ]
    fmts = {
        "Forecast accuracy (1 − WAPE)": "0.0%",
        "Order value": "€#,##0",
        "Weekly forecast (units)": "#,##0.0",
        "Safety stock (units)": "#,##0",
        "Reorder point (units)": "#,##0",
        "EOQ (units)": "#,##0",
    }
    for i, (label, value, note) in enumerate(lines):
        rr = 4 + i
        s.cell(rr, 1, label).font = BOLD
        c = s.cell(rr, 2, value)
        c.number_format = fmts.get(label, "General")
        c.alignment = WRAP
        if isinstance(value, str) and value.startswith("="):
            c.font, c.fill = LINK_FONT, RESULT_FILL
        if note:
            s.cell(rr, 3, note).font = SUBTITLE
    s.merge_cells(start_row=7, start_column=2, end_row=7, end_column=3)
    s.row_dimensions[7].height = 48
    guide_row = 4 + len(lines) + 1
    s.cell(guide_row, 1, "How to use this workbook").font = BOLD
    guide = [
        "Inputs: change any blue-on-yellow cell (service level, holding rate, lead time, stock, POs) and every sheet recalculates.",
        "Forecast: the four walk-forward forecasts and their errors, week by week; the evaluation window is shaded.",
        "Method_Selection: WAPE, bias and error SD per method and the selection rule.",
        "Replenishment: safety stock → reorder point → inventory position → EOQ → order → status.",
        "Projection: 12 weeks of opening + receipts − demand = closing. Check: Excel vs Python, cell by cell.",
        "A portfolio project by Syed Danish Ali · linkedin.com/in/syeddanishali16",
    ]
    for i, text in enumerate(guide):
        s.cell(guide_row + 1 + i, 1, text).font = SUBTITLE

    demand_chart = LineChart()
    demand_chart.title, demand_chart.height, demand_chart.width = (
        "Actual demand vs selected forecast (last 52 weeks)",
        7.5,
        22,
    )
    demand_chart.add_data(Reference(fs, min_col=3, min_row=4, max_row=4), titles_from_data=True)
    demand_chart.series[0].val.numRef.f = f"Forecast!$C${last - 51}:$C${last}"
    demand_chart.add_data(Reference(fs, min_col=16, min_row=4, max_row=4), titles_from_data=True)
    demand_chart.series[1].val.numRef.f = f"Forecast!$P${last - 51}:$P${last}"
    demand_chart.set_categories(Reference(fs, min_col=1, min_row=last - 51, max_row=last))
    s.add_chart(demand_chart, f"A{guide_row + len(guide) + 3}")
    return wb


def sku_workbook_bytes(result: PlanResult, sku: str) -> bytes:
    buffer = io.BytesIO()
    build_sku_workbook(result, sku).save(buffer)
    return buffer.getvalue()
