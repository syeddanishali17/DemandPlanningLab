"""Extract the workbook's CALCULATED results into CSV fixtures for the Excel-parity tests.

These fixtures are the "expected" side of ``tests/test_excel_parity.py``. They cover all
48 SKUs (forecast selection, safety stock, ROP, EOQ, recommended order, projection,
status) and the portfolio backtest results for both policies.

Usage (from the repository root)::

    python scripts/extract_excel_reference.py --workbook docs/validation/Demand_Inventory_Planning_Model_V1.xlsx
"""

from __future__ import annotations

import argparse
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.utils import column_index_from_string as col

ROOT = Path(__file__).resolve().parents[1]
FIXTURE_DIR = ROOT / "tests" / "fixtures" / "excel_reference"

FIRST_SKU_ROW, LAST_SKU_ROW = 7, 54  # rows of the 48 SKUs on the per-SKU sheets


def _values(ws, row: int, mapping: dict[str, str]) -> dict:
    return {name: ws.cell(row, col(letter)).value for name, letter in mapping.items()}


def extract(workbook: Path) -> None:
    wb = openpyxl.load_workbook(workbook, data_only=True)

    # --- Forecast sheet: per-method metrics and selection ------------------------------
    fc_map = {
        "sku": "A",
        "wape_naive": "I",
        "wape_ma4": "J",
        "wape_ma8": "K",
        "wape_ses": "L",
        "bias_naive": "M",
        "bias_ma4": "N",
        "bias_ma8": "O",
        "bias_ses": "P",
        "sd_naive": "Q",
        "sd_ma4": "R",
        "sd_ma8": "S",
        "sd_ses": "T",
        "fcst_naive": "U",
        "fcst_ma4": "V",
        "fcst_ma8": "W",
        "fcst_ses": "X",
        "selected_method": "AE",
        "weekly_forecast": "AF",
        "wape": "AG",
        "bias": "AH",
        "error_sd": "AI",
        "calib_selected_method": "AZ",
        "calib_error_sd": "BA",
    }
    ws = wb["Forecast"]
    forecast = pd.DataFrame([_values(ws, r, fc_map) for r in range(FIRST_SKU_ROW, LAST_SKU_ROW + 1)])
    portfolio_wape = ws["E3"].value

    # --- Replenishment_Plan sheet -------------------------------------------------------
    rp_map = {
        "sku": "A",
        "safety_stock": "L",
        "lead_time_demand": "M",
        "reorder_point": "N",
        "on_hand": "O",
        "on_order": "P",
        "backorders": "Q",
        "inventory_position": "R",
        "weeks_of_cover_ip": "S",
        "weeks_of_cover_on_hand": "T",
        "annual_demand": "U",
        "holding_cost_per_unit": "Y",
        "eoq": "Z",
        "eoq_rounded": "AB",
        "reorder_trigger": "AC",
        "recommended_order": "AD",
        "first_zero_week": "AE",
        "zero_within_lead_time": "AF",
        "status": "AG",
        "reason": "AH",
        "on_hand_value": "AI",
        "attention_rank": "AK",
    }
    ws = wb["Replenishment_Plan"]
    plan = pd.DataFrame([_values(ws, r, rp_map) for r in range(FIRST_SKU_ROW, LAST_SKU_ROW + 1)])

    # --- Projection sheet: 12 closing-inventory weeks per SKU ----------------------------
    ws = wb["Projection"]
    rows = []
    for r in range(FIRST_SKU_ROW, LAST_SKU_ROW + 1):
        sku = ws.cell(r, col("A")).value
        for w in range(12):
            rows.append(
                {
                    "sku": sku,
                    "week": w + 1,
                    "existing_po_receipt": ws.cell(r, col("I") + w).value,
                    "new_order_receipt": ws.cell(r, col("U") + w).value,
                    "forecast_demand": ws.cell(r, col("AG") + w).value,
                    "closing_inventory": ws.cell(r, col("AS") + w).value,
                }
            )
    projection = pd.DataFrame(rows)

    # --- Backtest: portfolio totals, per-SKU-week detail ---------------------------------
    ws = wb["Backtest"]
    totals = pd.DataFrame(
        {
            "metric": [
                "fill_rate",
                "stockout_weeks",
                "average_inventory_units",
                "number_of_orders",
                "average_inventory_value",
            ],
            "baseline": [ws["B14"].value, ws["B15"].value, ws["B16"].value, ws["B17"].value, ws["B18"].value],
            "proposed": [ws["C14"].value, ws["C15"].value, ws["C16"].value, ws["C17"].value, ws["C18"].value],
        }
    )
    bt_map = {
        "sku": "A",
        "week": "B",
        "actual_demand": "I",
        "a_trailing_avg": "J",
        "a_target": "K",
        "a_receipts": "L",
        "a_stock_start": "M",
        "a_on_order": "N",
        "a_inventory_position": "O",
        "a_order": "P",
        "a_sales": "Q",
        "a_lost": "R",
        "a_closing": "S",
        "b_forecast": "V",
        "b_error_sd": "W",
        "b_safety_stock": "X",
        "b_reorder_point": "Y",
        "b_eoq": "Z",
        "b_eoq_rounded": "AA",
        "b_receipts": "AB",
        "b_stock_start": "AC",
        "b_on_order": "AD",
        "b_inventory_position": "AE",
        "b_order": "AF",
        "b_sales": "AG",
        "b_lost": "AH",
        "b_closing": "AI",
    }
    detail = pd.DataFrame([_values(ws, r, bt_map) for r in range(53, 821)])

    FIXTURE_DIR.mkdir(parents=True, exist_ok=True)
    forecast.to_csv(FIXTURE_DIR / "forecast.csv", index=False)
    plan.to_csv(FIXTURE_DIR / "replenishment_plan.csv", index=False)
    projection.to_csv(FIXTURE_DIR / "projection.csv", index=False)
    totals.to_csv(FIXTURE_DIR / "backtest_totals.csv", index=False)
    detail.to_csv(FIXTURE_DIR / "backtest_detail.csv", index=False)
    (FIXTURE_DIR / "portfolio_wape.txt").write_text(repr(portfolio_wape), encoding="utf-8")

    print(f"forecast            : {len(forecast)} SKUs")
    print(f"replenishment_plan  : {len(plan)} SKUs")
    print(f"projection          : {len(projection)} SKU-weeks")
    print(f"backtest_totals     :\n{totals}")
    print(f"backtest_detail     : {len(detail)} SKU-weeks")
    print(f"portfolio_wape      : {portfolio_wape}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, required=True)
    args = parser.parse_args()
    extract(args.workbook.resolve())
